"""Deck tests: the PowerPoint is built from plain data, so no Spark is needed."""

from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE

from sales_insights.deck.deck import DeckData, KpiCard, build_deck


def sample(alerts=True, forecast=True) -> DeckData:
    return DeckData(
        title="Weekly sales insights",
        subtitle="Data up to Sunday 25 October 2026",
        cards=[
            KpiCard("Net revenue", "LKR 160.2M", "up 4.1% vs 1-25 Oct 2025", True),
            KpiCard("Units (cases)", "78,410", "down 2.0% vs 1-25 Oct 2025", False),
            KpiCard("Invoices", "1,402", "up 1.2% vs 1-25 Oct 2025", True),
            KpiCard("Gross margin %", "21.4%", "up 0.6 points vs 1-25 Oct 2025", True),
        ],
        cards_caption="Month to date (1-25 Oct 2026) compared with the same days last year (1-25 Oct 2025)",
        key_points=["Possible stock-out: MC600 WEEDICIDE 1LTR in Western.", "Revenue up 4.1%.", "Northern rose most."],
        trend=[(f"M{i}", 100 + i * 5.0) for i in range(13)],
        trend_caption="Net revenue per month, LKR millions",
        movers={
            "By province": [("Northern", 12.5), ("Southern", -4.2), ("Western", 3.1)],
            "By product group": [("WEEDICIDE", 6.0), ("INSECTICIDE", -1.5)],
        },
        movers_caption="Sep 2026 vs Sep 2025: change in net revenue, LKR millions",
        alerts=["Possible stock-out: MC600 WEEDICIDE 1LTR has had no sales in Western for 3 days."] if alerts else [],
        summary=["Point one.", "Point two."],
        about=["Data up to 25 Oct 2026."],
        outlook=[("Apr 26", 150.0, None), ("May 26", 160.0, None), ("Oct 26", None, 187.9), ("Nov 26", None, 184.0)]
        if forecast
        else [],
        outlook_caption="Net revenue per month, LKR millions",
        outlook_notes=["Oct 2026: about LKR 187.9M (likely LKR 175.0M to LKR 200.8M)"] if forecast else [],
    )


def _texts(slide):
    return " ".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)


def test_deck_has_eight_slides_with_native_charts(tmp_path):
    path = build_deck(sample(), tmp_path / "deck.pptx")
    prs = Presentation(path)
    assert len(prs.slides) == 8
    s = list(prs.slides)
    assert "Weekly sales insights" in _texts(s[0])
    assert "LKR 160.2M" in _texts(s[1]) and "What to know" in _texts(s[1])
    charts = [sh.chart for sh in s[2].shapes if sh.has_chart]
    assert len(charts) == 1 and charts[0].chart_type == XL_CHART_TYPE.LINE_MARKERS
    assert len(list(charts[0].plots[0].categories)) == 13
    [outlook] = [sh.chart for sh in s[3].shapes if sh.has_chart]
    assert outlook.chart_type == XL_CHART_TYPE.COLUMN_CLUSTERED
    assert [ser.name for ser in outlook.series] == ["Actual", "Forecast"]
    assert list(outlook.series[1].values) == [None, None, 187.9, 184.0]
    assert "likely LKR 175.0M" in _texts(s[3])
    bars = [sh.chart for sh in s[4].shapes if sh.has_chart]
    assert len(bars) == 2 and all(c.chart_type == XL_CHART_TYPE.BAR_CLUSTERED for c in bars)
    assert "MC600 WEEDICIDE 1LTR" in _texts(s[5])
    assert "Point two." in _texts(s[6])


def test_deck_without_alerts_says_so(tmp_path):
    prs = Presentation(build_deck(sample(alerts=False), tmp_path / "deck.pptx"))
    assert "No alerts this week" in _texts(list(prs.slides)[5])


def test_deck_without_forecast_says_so(tmp_path):
    prs = Presentation(build_deck(sample(forecast=False), tmp_path / "deck.pptx"))
    s = list(prs.slides)[3]
    assert "Not enough history for a forecast" in _texts(s) and not any(sh.has_chart for sh in s.shapes)


def test_stock_slide_appears_only_with_stock_findings(tmp_path):
    d = sample()
    d.stock = [
        "Confirmed stock-out: Colombo Depot has had no MC600 WEEDICIDE 1LTR (CASE) since 22 Oct (8 days).",
        "Low cover: Galle Depot has 1 MAN1KG left, about 0.4 days at its usual 2.5 a day.",
        "Overstock: 47 stock positions hold more than 120 days of cover, LKR 12.2M of stock above that level.",
    ]
    d.stock_caption = "Stock data up to 29 Oct 2026"
    prs = Presentation(str(build_deck(d, tmp_path / "with_stock.pptx")))
    assert len(prs.slides) == 9
    stock = list(prs.slides)[6]  # right after the alerts slide
    assert "Stock data up to 29 Oct 2026" in _texts(stock) and "Confirmed stock-out: Colombo Depot" in _texts(stock)
    assert "Summary" in _texts(list(prs.slides)[7])
