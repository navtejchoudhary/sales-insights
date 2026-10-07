"""Weekly sales deck (PowerPoint) built with python-pptx. Plain Python: no Spark in here.

The insights job fills a DeckData with numbers and sentences; build_deck() lays them out.
Charts are native PowerPoint charts, so anyone can click them, see values and restyle them.

    slide 1  title
    slide 2  at a glance: 4 KPI cards + the 3 most important points
    slide 3  revenue trend: last 13 complete months
    slide 4  what moved revenue: change by province and by product group
    slide 5  alerts: possible stock-outs, unusual days
    slide 6  summary: the narrative
    slide 7  about these numbers: definitions and data checks
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_TICK_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

FOREST = RGBColor(0x2C, 0x5F, 0x2D)  # dominant: crop-protection green
MOSS = RGBColor(0x97, 0xBC, 0x62)
MIST = RGBColor(0xEE, 0xF3, 0xE8)  # card background
INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x5B, 0x65, 0x70)
UP = RGBColor(0x2E, 0x7D, 0x32)
DOWN = RGBColor(0xB2, 0x3A, 0x48)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HEAD = "Cambria"
BODY = "Calibri"

W, H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.6)


@dataclass
class KpiCard:
    label: str
    value: str
    change: str  # e.g. "up 4.1% vs 1-5 Sep 2026"
    good: bool | None = None  # True green, False red, None neutral


@dataclass
class DeckData:
    title: str
    subtitle: str
    cards: list[KpiCard]
    cards_caption: str
    key_points: list[str]
    trend: list[tuple[str, float]]  # (month label, net revenue in LKR millions)
    trend_caption: str
    movers: dict[str, list[tuple[str, float]]]  # chart title -> [(item, change in LKR millions)]
    movers_caption: str
    alerts: list[str]
    summary: list[str]
    about: list[str] = field(default_factory=list)


# ---------------------------------------------------------------------------
# Small drawing helpers
# ---------------------------------------------------------------------------


def _text(
    slide, x, y, w, h, text, size=16, bold=False, color=INK, font=BODY, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP
):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, 0)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size, run.font.bold, run.font.name = Pt(size), bold, font
    run.font.color.rgb = color
    return box


def _bullets(slide, x, y, w, h, items, size=16, color=INK, space=10):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, 0)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(space)
        dot = p.add_run()
        dot.text = "●  "
        dot.font.size, dot.font.color.rgb, dot.font.name = Pt(size - 4), MOSS, BODY
        run = p.add_run()
        run.text = item
        run.font.size, run.font.color.rgb, run.font.name = Pt(size), color, BODY
    return box


def _title(slide, text, caption=None):
    _text(slide, MARGIN, Inches(0.45), W - 2 * MARGIN, Inches(0.8), text, size=34, bold=True, color=FOREST, font=HEAD)
    if caption:
        _text(slide, MARGIN, Inches(1.2), W - 2 * MARGIN, Inches(0.4), caption, size=14, color=MUTED)


def _card(slide, x, y, w, h, fill=MIST):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.adjustments[0] = 0.08
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def _blank(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE
    return slide


def _style_axes(chart, number_format):
    va = chart.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = RGBColor(0xE5, 0xE7, 0xEB)
    va.format.line.fill.background()
    va.tick_labels.font.size, va.tick_labels.font.color.rgb = Pt(11), MUTED
    va.tick_labels.number_format, va.tick_labels.number_format_is_linked = number_format, False
    ca = chart.category_axis
    ca.tick_labels.font.size, ca.tick_labels.font.color.rgb = Pt(11), MUTED
    ca.format.line.color.rgb = RGBColor(0xD1, 0xD5, 0xDB)


# ---------------------------------------------------------------------------
# Slides
# ---------------------------------------------------------------------------


def _slide_title(prs, d: DeckData):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = FOREST
    _text(s, MARGIN, Inches(2.4), Inches(11), Inches(1.4), d.title, size=48, bold=True, color=WHITE, font=HEAD)
    _text(s, MARGIN, Inches(3.8), Inches(11), Inches(0.6), d.subtitle, size=20, color=MOSS)
    _text(s, MARGIN, Inches(6.6), Inches(11), Inches(0.4),
          "Generated automatically from the sales lakehouse. Money in LKR, excluding VAT.", size=12, color=MIST)  # fmt: skip


def _slide_glance(prs, d: DeckData):
    s = _blank(prs)
    _title(s, "At a glance", d.cards_caption)
    n = max(len(d.cards), 1)
    gap = Inches(0.3)
    cw = int((W - 2 * MARGIN - gap * (n - 1)) / n)
    y, ch = Inches(1.85), Inches(1.9)
    for i, c in enumerate(d.cards):
        x = MARGIN + i * (cw + gap)
        _card(s, x, y, cw, ch)
        pad = Inches(0.25)
        _text(s, x + pad, y + Inches(0.2), cw - 2 * pad, Inches(0.35), c.label, size=13, color=MUTED)
        _text(
            s, x + pad, y + Inches(0.55), cw - 2 * pad, Inches(0.75), c.value, size=28, bold=True, color=INK, font=HEAD
        )
        color = UP if c.good is True else DOWN if c.good is False else MUTED
        _text(s, x + pad, y + Inches(1.35), cw - 2 * pad, Inches(0.45), c.change, size=12, color=color)
    _text(s, MARGIN, Inches(4.15), Inches(6), Inches(0.4), "What to know", size=20, bold=True, color=FOREST, font=HEAD)
    _bullets(s, MARGIN, Inches(4.65), W - 2 * MARGIN, Inches(2.5), d.key_points[:3], size=15)


def _slide_trend(prs, d: DeckData):
    s = _blank(prs)
    _title(s, "Revenue trend", d.trend_caption)
    cd = CategoryChartData()
    cd.categories = [m for m, _ in d.trend]
    cd.add_series("Net revenue (LKR M)", [round(v, 1) for _, v in d.trend])
    gf = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, MARGIN, Inches(1.8), W - 2 * MARGIN, Inches(5.2), cd)
    chart = gf.chart
    chart.has_title = False
    chart.has_legend = False
    chart.font.name = BODY
    ser = chart.series[0]
    ser.format.line.color.rgb, ser.format.line.width = FOREST, Pt(2.75)
    ser.smooth = False
    ser.marker.format.fill.solid()
    ser.marker.format.fill.fore_color.rgb = FOREST
    ser.marker.format.line.color.rgb = FOREST
    plot = chart.plots[0]
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format, dl.number_format_is_linked = '#,##0"M"', False
    dl.position = XL_LABEL_POSITION.ABOVE
    dl.font.size, dl.font.color.rgb = Pt(10), MUTED
    _style_axes(chart, '#,##0"M"')


def _slide_movers(prs, d: DeckData):
    s = _blank(prs)
    _title(s, "What moved revenue", d.movers_caption)
    items = list(d.movers.items())[:2]
    gap = Inches(0.4)
    cw = int((W - 2 * MARGIN - gap) / 2)
    for i, (title, rows) in enumerate(items):
        x = MARGIN + i * (cw + gap)
        _text(s, x, Inches(1.8), cw, Inches(0.4), title, size=16, bold=True, color=INK, font=HEAD)
        if not rows:
            _text(s, x, Inches(2.4), cw, Inches(0.5), "No change above the noise level.", size=14, color=MUTED)
            continue
        rows = sorted(rows, key=lambda r: r[1])
        cd = CategoryChartData()
        cd.categories = [r[0] for r in rows]
        cd.add_series("Change (LKR M)", [round(r[1], 1) for r in rows])
        chart = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, x, Inches(2.25), cw, Inches(4.8), cd).chart
        chart.has_title = False
        chart.has_legend = False
        chart.font.name = BODY
        plot = chart.plots[0]
        plot.gap_width = 60
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.number_format, dl.number_format_is_linked = '+#,##0.0"M";-#,##0.0"M"', False
        dl.position = XL_LABEL_POSITION.OUTSIDE_END
        dl.font.size, dl.font.color.rgb = Pt(10), INK
        ser = chart.series[0]
        ser.invert_if_negative = False
        for j, (_, v) in enumerate(rows):
            pt = ser.points[j]
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = UP if v >= 0 else DOWN
        _style_axes(chart, '#,##0"M"')
        chart.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW  # labels left of the bars


def _slide_alerts(prs, d: DeckData):
    s = _blank(prs)
    _title(
        s, "Alerts", "Checked every day: products that stopped selling where they usually sell daily, and unusual days"
    )
    if not d.alerts:
        _card(s, MARGIN, Inches(2.0), W - 2 * MARGIN, Inches(1.4))
        _text(s, MARGIN + Inches(0.4), Inches(2.0), W - 2 * MARGIN - Inches(0.8), Inches(1.4),
              "No alerts this week: every regular product is still selling in its usual provinces, and no day was unusual.",
              size=18, color=INK, anchor=MSO_ANCHOR.MIDDLE)  # fmt: skip
        return
    y = Inches(1.9)
    for a in d.alerts[:4]:
        _card(s, MARGIN, y, W - 2 * MARGIN, Inches(1.05))
        badge = s.shapes.add_shape(MSO_SHAPE.OVAL, MARGIN + Inches(0.3), y + Inches(0.3), Inches(0.45), Inches(0.45))
        badge.fill.solid()
        badge.fill.fore_color.rgb = DOWN
        badge.line.fill.background()
        _text(s, MARGIN + Inches(0.3), y + Inches(0.3), Inches(0.45), Inches(0.45), "!", size=18, bold=True,
              color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)  # fmt: skip
        _text(s, MARGIN + Inches(1.0), y, W - 2 * MARGIN - Inches(1.3), Inches(1.05), a, size=15, color=INK,
              anchor=MSO_ANCHOR.MIDDLE)  # fmt: skip
        y += Inches(1.25)


def _slide_summary(prs, d: DeckData):
    s = _blank(prs)
    _title(s, "Summary", "The most important points, in order")
    points = d.summary[:6]
    if not points:
        return
    gap = Inches(0.18)
    ch = min(Inches(0.95), int((Inches(5.3) - gap * (len(points) - 1)) / len(points)))
    y = Inches(1.85)
    for i, text in enumerate(points, start=1):
        _card(s, MARGIN, y, W - 2 * MARGIN, ch)
        dot = s.shapes.add_shape(
            MSO_SHAPE.OVAL, MARGIN + Inches(0.25), y + (ch - Inches(0.5)) // 2, Inches(0.5), Inches(0.5)
        )
        dot.fill.solid()
        dot.fill.fore_color.rgb = FOREST
        dot.line.fill.background()
        _text(s, MARGIN + Inches(0.25), y + (ch - Inches(0.5)) // 2, Inches(0.5), Inches(0.5), str(i), size=16, bold=True,
              color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)  # fmt: skip
        _text(s, MARGIN + Inches(1.0), y, W - 2 * MARGIN - Inches(1.3), ch, text, size=15, color=INK,
              anchor=MSO_ANCHOR.MIDDLE)  # fmt: skip
        y += ch + gap


def _slide_about(prs, d: DeckData):
    s = _blank(prs)
    _title(s, "About these numbers")
    _bullets(s, MARGIN, Inches(1.5), W - 2 * MARGIN, Inches(5.6), d.about, size=14, color=MUTED, space=8)


def build_deck(d: DeckData, path: str | Path) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    for make in (_slide_title, _slide_glance, _slide_trend, _slide_movers, _slide_alerts, _slide_summary, _slide_about):
        make(prs, d)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    # Build the .pptx (a zip file) in memory, then write it in one go: Unity Catalog volumes refuse
    # the random writes a zip writer makes when saving straight to /Volumes/...
    buf = io.BytesIO()
    prs.save(buf)
    path.write_bytes(buf.getvalue())
    return path
